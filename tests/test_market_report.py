"""market_report 시장 섹션 + weekly_report docx 표 분기 테스트. API 호출 없음."""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import weekly_report  # noqa: E402
from market_context import WATCH  # noqa: E402
from market_report import build_market_section  # noqa: E402

CTX = ROOT / "data" / "market" / "context_latest.json"


def _rows(md):
    return [ln for ln in md.splitlines() if ln.startswith("|") and not ln.startswith("|---")][1:]   # 헤더 제외


def test_real_context_seven_rows_and_low_volume():
    md = build_market_section(str(CTX))
    rows = _rows(md)
    assert len(rows) == 7
    como = next(r for r in rows if r.startswith("| COMO"))
    assert [c.strip() for c in como.strip("|").split("|")][1:] == ["—", "—", "—", "판정 제외(물량 부족)", "—"]
    assert "기준월 " in md and "출처 Census BPS·SOC" in md and "미착공·취소 미반영" in md


def test_missing_file_no_exception(tmp_path):
    md = build_market_section(str(tmp_path / "nope.json"))
    assert "시장 데이터 없음" in md and "|" not in md


def test_broken_json_no_exception(tmp_path):
    p = tmp_path / "broken.json"
    p.write_text('{"as_of": "2026-08", "cbsa": {', encoding="utf-8")
    md = build_market_section(str(p))
    assert "시장 데이터 없음" in md and "|" not in md


def test_no_watch_key_falls_back_to_constant(tmp_path):
    d = json.loads(CTX.read_text(encoding="utf-8"))
    d.pop("watch", None)
    p = tmp_path / "nowatch.json"
    p.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    rows = _rows(build_market_section(str(p)))
    assert [r.split("|")[1].strip().split(" · ")[0] for r in rows] == [s for _, s in WATCH]


def test_insert_before_implications_or_append():
    sec = "## 시장 데이터(공급) — WATCH 7\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    with_h = "## 1. 개발 현황\n- x\n\n## 4. 시사점 (우미글로벌 관점)\n- y"
    out = weekly_report.insert_market_section(with_h, sec)
    assert out.index("## 시장 데이터") < out.index("## 4. 시사점") and out.endswith("- y")
    without_h = "## 1. 개발 현황\n- x"
    out2 = weekly_report.insert_market_section(without_h, sec)
    assert out2.startswith(without_h) and out2.rstrip().endswith("| 1 | 2 |")


def test_reference_block_labels_only(monkeypatch):
    import market_report
    monkeypatch.setattr(market_report, "build_market_labels", lambda *a, **k: None)
    assert weekly_report.market_reference_block() == ""
    monkeypatch.undo()
    blk = weekly_report.market_reference_block()   # 실제 context_latest.json
    assert "--- 참고: 지역 공급 판정(자동 규칙, 참고용) ---" in blk
    assert ("ATL 평균 수준 · DFW 평균 수준 · LA 허가 급증 · SEA 허가 급증 · DEN 평균 수준 · "
            "PHX 평균 수준 · COMO 판정 제외") in blk
    assert "(ATL·평균 수준)" in blk and "판정 제외 지역은 라벨을 붙이지 말라" in blk
    assert "|" not in blk and "12,140" not in blk and "9,878" not in blk   # 표·수치·준공 추정치 미전달


def test_index_html_labels_match_python():
    """index.html 내장 LABELS_KO 와 market_rules.LABELS_KO(정본)가 어긋나면 실패."""
    import re
    from market_rules import LABELS_KO
    html = (ROOT / "index.html").read_text(encoding="utf-8")
    m = re.search(r"const LABELS_KO = (\{[^}]*\});", html)
    assert m, "index.html에 const LABELS_KO = {...}; 한 줄이 있어야 한다"
    assert json.loads(m.group(1)) == LABELS_KO
    assert weekly_report.MARKET_LEAK_TERMS[:4] == tuple(LABELS_KO.values())


def test_leakage_flags_implications_only():
    md = ("## 1. 개발 현황\n- Hines가 Buckhead 312유닛 매입(ATL·평균 수준). LA 착공(LA·허가 급증).\n\n"
          "## 시장 데이터(공급) — WATCH 7\n| 지역 | 판정 |\n|---|---|\n| ATL | 평균 수준 |\n\n"
          "## 4. 시사점 (우미글로벌 관점)\n"
          "- PHX는 평균 수준 국면으로 BTR 여건 양호. JV 기회 주목.\n"
          "- DFW는 2027 예상 준공 물량이 높아 모니터링 필요.\n\n---\n*맺음말*")
    hits = weekly_report.check_market_leakage(md)
    assert hits == ["PHX는 평균 수준 국면으로 BTR 여건 양호.", "DFW는 2027 예상 준공 물량이 높아 모니터링 필요."]


def test_leakage_clean_returns_empty():
    md = ("## 1. 개발 현황\n- 매입(ATL·평균 수준), 착공(SEA·허가 급증).\n\n"
          "## 4. 시사점\n- 일본계 자본 JV 진입 패턴 정착. 디스트레스 자산 매입 검토.\n"
          "- Village at Sonoran Vista가 신규 준공 후 리파이낸스 확보.\n")
    assert weekly_report.check_market_leakage(md) == []   # "신규 준공 후"는 기사 사실 서술 → 통과


def test_leakage_spacing_variants():
    md = ("## 1. 개발 현황\n- 소형 자산 거래(LA·허가급증).\n\n"
          "## 4. 시사점\n- LA는 허가급증 시장. SEA는 허가  급증 국면. DFW는 판정제외 지역. JV 기회.\n")
    assert weekly_report.check_market_leakage(md) == ["LA는 허가급증 시장.", "SEA는 허가  급증 국면.", "DFW는 판정제외 지역."]


def test_leakage_completion_terms():
    md = "## 4. 시사점\n- DFW는 2027 예상 준공 물량이 높아 모니터링 필요.\n- 신규 준공 후 리파이낸스 확보.\n"
    assert weekly_report.check_market_leakage(md) == ["DFW는 2027 예상 준공 물량이 높아 모니터링 필요."]


def test_region_tokens_missing_db(tmp_path):
    assert weekly_report.load_watch_regions(["a1"], db_path=str(tmp_path / "none.db")) == {}
    arts = [{"article_id": "a1", "title": "T", "summary": "S", "event_tags": "x"}]
    assert weekly_report.format_articles(arts) == weekly_report.format_articles(arts, {})
    assert weekly_report.format_articles(arts, {"a1": "LA"}).endswith("태그: x [지역: LA]")


def test_docx_table_branch_and_fallback(tmp_path, monkeypatch):
    from docx import Document
    monkeypatch.setattr(weekly_report, "REPORTS_DIR", str(tmp_path))
    md = ("## 섹션\n\n일반 문단입니다.\n\n"
          "| 지역 | 비율 |\n|---|---|\n| ATL | 0.91 |\n\n"
          "| 깨진 | 표 |\n| 셀 하나 |\n")
    path = weekly_report.save_report_docx(md, "weekly")
    doc = Document(path)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert (len(t.rows), len(t.columns)) == (2, 2)
    assert t.cell(0, 0).text == "지역" and t.cell(1, 1).text == "0.91"
    assert t.cell(0, 0).paragraphs[0].runs[0].bold is True
    texts = [p.text for p in doc.paragraphs]
    assert "일반 문단입니다." in texts
    assert "| 깨진 | 표 |" in texts and "| 셀 하나 |" in texts   # 셀 수 불일치 → 텍스트 폴백
